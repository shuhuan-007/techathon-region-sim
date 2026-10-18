# -*- coding: utf-8 -*-
"""Verify a merged DOCX: heading sequence, table widths, stray H1s."""
import re
import sys

from docx import Document

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
CJK = "一二三四五六七八九十"


def cjk_value(s):
    s = s.strip()
    if "十" not in s:
        return CJK.index(s) + 1
    left, _, right = s.partition("十")
    tens = CJK.index(left) + 1 if left else 1
    ones = CJK.index(right) + 1 if right else 0
    return tens * 10 + ones


def cjk_numeral(n):
    if n <= 10:
        return CJK[n - 1]
    if n < 20:
        return "十" + CJK[n - 11]
    tens, ones = divmod(n, 10)
    return CJK[tens - 1] + "十" + (CJK[ones - 1] if ones else "")


doc = Document(sys.argv[1])
sec = doc.sections[0]
usable = (sec.page_width - sec.left_margin - sec.right_margin) / 360000

h1 = h2 = h3 = 0
expect = 1
problems = []
for p in doc.paragraphs:
    st = p.style.name
    txt = p.text.strip()
    if st == "Heading 1":
        h1 += 1
        if txt and not txt.startswith(tuple("0123456789")):
            print(f"  H1: {txt}")
    elif st == "Heading 2":
        h2 += 1
        m = re.match(r"^([一二三四五六七八九十]+)、", txt)
        if not m:
            problems.append(f"H2 unnumbered: {txt[:40]}")
        else:
            val = cjk_value(m.group(1))
            if val != expect:
                problems.append(f"H2 sequence broken at expected {expect}: got {txt[:30]}")
            expect = val + 1
    elif st == "Heading 3":
        h3 += 1
        m = re.match(r"^(\d+)\.(\d+)\s", txt)
        if not m:
            problems.append(f"H3 unnumbered: {txt[:40]}")

# table widths
narrow = 0
bad_sum = 0
for i, t in enumerate(doc.tables, 1):
    g = [int(x.get(W + "w")) / 567 for x in t._tbl.tblGrid.findall(W + "gridCol")]
    if abs(sum(g) - usable) > 0.08:
        bad_sum += 1
        problems.append(f"T{i} width {round(sum(g),2)} != {round(usable,2)}")
    if any(x < 0.75 for x in g):
        narrow += 1
        problems.append(f"T{i} narrow col {round(min(g),2)}")

print(f"headings: H1={h1} H2={h2} H3={h3} | chapters expected end at {cjk_numeral(expect-1)}")
print(f"tables={len(doc.tables)} width_mismatch={bad_sum} narrow={narrow}")
print("problems:", problems if problems else "none")
